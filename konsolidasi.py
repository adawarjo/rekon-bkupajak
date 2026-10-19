"""Rekonsiliasi Pajak BKU SKPD SIPD + Coretax: mencocokkan BKU SKPD dengan Buku Besar Coretax.

Dipakai oleh app.py (GUI), atau langsung dari command line:
    python konsolidasi.py "BKU SKPD.xlsx" bukubesar_coretax.csv hasil.xlsx
"""
from __future__ import annotations

import csv
import dataclasses
import datetime as dt
import io
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

KAP_DEPOSIT = "411618"
MAKS_SALAH_KETIK = 3  # selisih karakter NTPN maksimal yang masih dianggap salah ketik
POLA_ILMIAH = re.compile(r"-?\d+([.,]\d+)?E[+-]?\d+", re.I)  # angka panjang yang dirusak Excel, mis. 3,11661E+12

COCOK = "COCOK"
TANGGAL_BEDA = "COCOK - TANGGAL BEDA"
NIHIL = "NIHIL"
SALAH_KETIK = "NTPN SALAH KETIK"
NILAI_BEDA = "NILAI BEDA"
TIDAK_DITEMUKAN = "TIDAK DITEMUKAN"
NTPN_GANDA = "NTPN GANDA DI BKU"
DIBATALKAN = "NTPN DIBATALKAN DI CORETAX"

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]


# --------------------------------------------------------------------------- utilitas

def parse_angka(v) -> float:
    """Ubah angka format Indonesia ('2.171.379', '1.234,50') atau angka biasa menjadi float."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace("Rp", "").replace(" ", "")
    if s in ("", "-"):
        return 0.0
    negatif = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "." in s:
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
    elif "," in s:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        raise ValueError(f"Nilai '{v}' bukan angka yang valid") from None
    return round(-x if negatif else x, 2)


def parse_tanggal(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # nomor seri tanggal Excel
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    s = str(v).strip()
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return dt.date(*map(int, m.groups()))
    m = re.match(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", s)
    if m:
        d, mo, y = map(int, m.groups())
        return dt.date(y + 2000 if y < 100 else y, mo, d)
    return None


def norm_ntpn(v) -> str:
    s = re.sub(r"\s+", "", str(v if v is not None else "")).upper()
    return "" if s in ("-", "NONE", "NAN", "0") else s


def jarak_edit(a: str, b: str) -> int:
    """Jarak Levenshtein: jumlah karakter yang harus diubah/ditambah/dihapus."""
    if a == b:
        return 0
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def rupiah(x: float) -> str:
    s = f"{abs(x):,.0f}" if round(x) == x else f"{abs(x):,.2f}"
    s = s.replace(",", "#").replace(".", ",").replace("#", ".")
    return ("-Rp" if x < 0 else "Rp") + s


def tgl(d) -> str:
    return d.strftime("%d/%m/%Y") if isinstance(d, dt.date) else (d or "-")


def _sama(a: float, b: float) -> bool:
    return abs(a - b) < 0.005


def _teks_kode(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _ke_int(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


# --------------------------------------------------------------------------- baca file

def _baca_csv(path: Path) -> list[list[str]]:
    csv.field_size_limit(2**31 - 1)
    for enc in ("utf-8-sig", "cp1252"):
        try:
            teks = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    baris1 = teks.split("\n", 1)[0]
    pemisah = max(";,\t|", key=baris1.count)
    return list(csv.reader(io.StringIO(teks), delimiter=pemisah))


def _baca_sheet(path: Path) -> list[list[list]]:
    """Kembalikan isi file sebagai daftar sheet; tiap sheet berupa daftar baris."""
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        finally:
            wb.close()
    if ext in (".csv", ".txt"):
        return [_baca_csv(path)]
    if ext == ".xls":
        raise ValueError(f"{path.name}: format .xls lama belum didukung. Buka di Excel lalu Save As .xlsx.")
    raise ValueError(f"{path.name}: format file {ext} tidak didukung.")


def _norm_kolom(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip().lower()


def _cari_header(sheets, wajib):
    for rows in sheets:
        for i, r in enumerate(rows[:30]):
            nama = [_norm_kolom(c) for c in r]
            if all(w in nama for w in wajib):
                kol = {}
                for j, n in enumerate(nama):
                    if n:
                        kol.setdefault(n, j)
                return rows, i, kol
    return None


def _ambil(r, kol, nama):
    j = kol.get(nama)
    return r[j] if j is not None and j < len(r) else None


@dataclass
class BarisBku:
    no: int
    tanggal: dt.date | None
    tanggal_teks: str
    bukti: str
    uraian: str
    jenis: str  # "terima" atau "bayar"
    pajak: str
    kode_rek: str
    billing: str
    ntpn_asli: str
    ntpn: str
    terima: float
    bayar: float

    @property
    def bukti6(self) -> str:
        """Nomor urut 6 digit dari Nomor Bukti, mis. '33.26/06.0/002556/UP/...' -> '002556'."""
        m = re.search(r"(?:^|/)(\d{6})(?:/|$)", self.bukti) or re.search(r"(?<!\d)(\d{6})(?!\d)", self.bukti)
        return m.group(1) if m else ""


def _awalan_bukti(b: BarisBku | None) -> str:
    return f"Bukti {b.bukti6} : " if b and b.bukti6 else ""


@dataclass
class BarisCoretax:
    ntpn: str
    amount: float
    sisa: float | None
    kap: str
    kjs: str
    masa: str
    tanggal: dt.date | None
    posting: str
    batal: bool
    jumlah_baris: int = 1

    @property
    def kap_kjs(self) -> str:
        return f"{self.kap}-{self.kjs}" if self.kjs else self.kap


def baca_bku(path) -> list[BarisBku]:
    path = Path(path)
    ketemu = _cari_header(_baca_sheet(path), ("ntpn", "uraian", "penerimaan", "pengeluaran"))
    if not ketemu:
        raise ValueError(
            f"{path.name} tidak dikenali sebagai BKU SKPD.\n"
            "Kolom NTPN, Uraian, Penerimaan dan Pengeluaran tidak ditemukan.")
    rows, hi, kol = ketemu
    hasil = []
    for idx, r in enumerate(rows[hi + 1:], start=hi + 2):
        if "no" in kol:
            no = _ke_int(_ambil(r, kol, "no"))
            if no is None:  # baris "Saldo Sebelumnya", "Total", dsb.
                continue
        else:
            no = idx
        uraian = str(_ambil(r, kol, "uraian") or "").strip()
        terima = parse_angka(_ambil(r, kol, "penerimaan"))
        bayar = parse_angka(_ambil(r, kol, "pengeluaran"))
        u = uraian.lower()
        if u.startswith("penerimaan"):
            jenis = "terima"
        elif u.startswith("pembayaran"):
            jenis = "bayar"
        elif terima:
            jenis = "terima"
        elif bayar:
            jenis = "bayar"
        else:
            continue
        if ":" in uraian and re.match(r"(penerimaan|pembayaran)\b", u):
            uraian = uraian.split(":", 1)[1].strip()
        billing_v = _ambil(r, kol, "id billing")
        billing = _teks_kode(billing_v)
        if isinstance(billing_v, (int, float)) and len(billing) < 15:
            billing = billing.zfill(15)  # nol di depan hilang karena tersimpan sebagai angka
        tgl_v = _ambil(r, kol, "tanggal")
        ntpn_asli = str(_ambil(r, kol, "ntpn") or "").strip()
        hasil.append(BarisBku(
            no=no, tanggal=parse_tanggal(tgl_v), tanggal_teks=str(tgl_v or ""),
            bukti=str(_ambil(r, kol, "nomor bukti") or "").strip(), uraian=uraian, jenis=jenis,
            pajak=str(_ambil(r, kol, "nama rekening") or "").strip(),
            kode_rek=_teks_kode(_ambil(r, kol, "kode rekening")), billing=billing,
            ntpn_asli=ntpn_asli, ntpn=norm_ntpn(ntpn_asli), terima=terima, bayar=bayar))
    if not hasil:
        raise ValueError(f"{path.name}: tidak ada baris transaksi pajak di BKU.")
    return hasil


def baca_coretax(path):
    path = Path(path)
    ketemu = _cari_header(_baca_sheet(path), ("referencenumber", "amount"))
    if not ketemu:
        raise ValueError(
            f"{path.name} tidak dikenali sebagai Buku Besar Coretax.\n"
            "Kolom ReferenceNumber dan Amount tidak ditemukan.")
    rows, hi, kol = ketemu
    data, dilewati, peringatan = [], 0, []
    wp_nama = wp_npwp = ""
    ntpn_rusak, npwp_rusak = 0, False
    for r in rows[hi + 1:]:
        ntpn = norm_ntpn(_ambil(r, kol, "referencenumber"))
        if not ntpn:
            continue
        if POLA_ILMIAH.fullmatch(ntpn):
            ntpn_rusak += 1
        jenis_ref = _norm_kolom(_ambil(r, kol, "referencetype"))
        if jenis_ref and jenis_ref != "payment":  # hanya setoran yang punya NTPN
            dilewati += 1
            continue
        sisa_v = _ambil(r, kol, "amountleft")
        data.append(BarisCoretax(
            ntpn=ntpn, amount=parse_angka(_ambil(r, kol, "amount")),
            sisa=parse_angka(sisa_v) if sisa_v not in (None, "") else None,
            kap=_teks_kode(_ambil(r, kol, "revenuecode")), kjs=_teks_kode(_ambil(r, kol, "paymentcode")),
            masa=_teks_kode(_ambil(r, kol, "periodcode")),
            tanggal=parse_tanggal(_ambil(r, kol, "transactiondate")) or parse_tanggal(_ambil(r, kol, "postingdate")),
            posting=str(_ambil(r, kol, "postingdate") or "")[:19].replace("T", " "),
            batal=_norm_kolom(_ambil(r, kol, "iscancelled")) in ("true", "1", "ya", "yes")))
        wp_nama = wp_nama or str(_ambil(r, kol, "taxpayername") or "").strip()
        npwp = _teks_kode(_ambil(r, kol, "taxpayertin"))
        if POLA_ILMIAH.fullmatch(npwp):
            npwp_rusak = True
        elif not wp_npwp:
            wp_npwp = npwp
    if not data:
        raise ValueError(f"{path.name}: tidak ada data setoran (NTPN) di Buku Besar Coretax.")
    if ntpn_rusak or npwp_rusak:
        peringatan.append(
            "PERHATIAN: file Coretax tampaknya pernah dibuka lalu disimpan ulang di Excel, sehingga angka panjang "
            "berubah ke format ilmiah (mis. NPWP 3,11661E+12)."
            + (f" {ntpn_rusak} NTPN ikut rusak sehingga tidak bisa dicocokkan." if ntpn_rusak else
               " NTPN dan Amount tetap utuh, tetapi sebaiknya pakai CSV asli unduhan Coretax.")
            + " Untuk melihat CSV, jangan klik Save di Excel.")
    return data, wp_nama, wp_npwp, dilewati, peringatan


# --------------------------------------------------------------------------- rekonsiliasi

@dataclass
class BarisRekon:
    terima: BarisBku
    bayar: BarisBku | None = None
    ct: BarisCoretax | None = None
    status: str = ""
    tindakan: list[str] = field(default_factory=list)

    @property
    def selisih(self) -> float | None:
        return None if self.ct is None else round(self.terima.terima - self.ct.amount, 2)

    @property
    def baris_bku(self) -> str:
        return f"No {self.terima.no}" + (f" & No {self.bayar.no}" if self.bayar else "")

    @property
    def tanggal_beda(self) -> bool:
        return bool(self.ct and self.ct.tanggal and self.terima.tanggal and self.ct.tanggal != self.terima.tanggal)


@dataclass
class Hasil:
    rekon: list[BarisRekon]
    coretax_valid: list[BarisCoretax]
    coretax_saja: list[BarisCoretax]
    coretax_batal: list[BarisCoretax]
    masalah_bku: list[str]
    catatan: list[str]
    jumlah_terima: int
    jumlah_bayar: int
    total_terima: float
    total_bayar: float
    total_coretax: float
    wp_nama: str
    wp_npwp: str
    masa: str
    file_bku: str
    file_coretax: str
    dibuat: dt.datetime

    @property
    def perlu_perbaikan(self) -> list[BarisRekon]:
        return [r for r in self.rekon if r.status not in (COCOK, NIHIL)]

    @property
    def jumlah_temuan(self) -> int:
        return len(self.perlu_perbaikan) + len(self.coretax_saja) + len(self.masalah_bku)


def _label_masa(tanggal: list[dt.date]) -> str:
    if not tanggal:
        return "-"
    bulan = {(d.year, d.month) for d in tanggal}
    if len(bulan) == 1:
        y, m = bulan.pop()
        return f"{BULAN[m - 1]} {y}"
    return f"{tgl(min(tanggal))} s.d. {tgl(max(tanggal))}"


def _pasangkan(rekon: list[BarisRekon], bayar: list[BarisBku]) -> list[str]:
    """Pasangkan tiap Penerimaan Pot. Pajak dengan baris Pembayaran-nya, laporkan yang janggal."""
    masalah, sisa = [], list(bayar)
    for r in rekon:
        t = r.terima
        calon = [b for b in sisa if b.ntpn == t.ntpn]
        if calon:
            calon.sort(key=lambda b: (not _sama(b.bayar, t.terima), b.billing != t.billing, b.no < t.no, abs(b.no - t.no)))
            r.bayar = calon[0]
            sisa.remove(r.bayar)
            if not _sama(r.bayar.bayar, t.terima):
                masalah.append(f"{_awalan_bukti(t)}BKU {r.baris_bku} (NTPN {t.ntpn_asli}): penerimaan {rupiah(t.terima)} "
                               f"tidak sama dengan pembayaran {rupiah(r.bayar.bayar)}")
    for r in rekon:
        if r.bayar:
            continue
        t = r.terima
        calon = [b for b in sisa if _sama(b.bayar, t.terima) and t.billing and b.billing == t.billing]
        if calon:
            r.bayar = min(calon, key=lambda b: abs(b.no - t.no))
            sisa.remove(r.bayar)
            masalah.append(f"{_awalan_bukti(t)}BKU {r.baris_bku}: NTPN baris penerimaan '{t.ntpn_asli}' berbeda dengan "
                           f"baris pembayaran '{r.bayar.ntpn_asli}'")
        elif t.terima or t.ntpn:
            masalah.append(f"{_awalan_bukti(t)}BKU No {t.no} (NTPN {t.ntpn_asli or '-'}, {rupiah(t.terima)}): "
                           "Penerimaan Pot. Pajak belum ada baris Pembayaran Pot. Pajak-nya")
    for b in sisa:
        if b.bayar or b.ntpn:
            masalah.append(f"{_awalan_bukti(b)}BKU No {b.no} (NTPN {b.ntpn_asli or '-'}, {rupiah(b.bayar)}): "
                           "Pembayaran Pot. Pajak tanpa baris Penerimaan pasangannya")
    return masalah


def rekonsiliasi(bku: list[BarisBku], coretax: list[BarisCoretax]):
    terima = [b for b in bku if b.jenis == "terima"]
    bayar = [b for b in bku if b.jenis == "bayar"]
    rekon = [BarisRekon(t) for t in terima]
    masalah_bku = _pasangkan(rekon, bayar)
    catatan = []

    # NTPN valid = tidak dibatalkan; NTPN yang muncul lebih dari sekali dijumlahkan
    valid: dict[str, BarisCoretax] = {}
    batal: dict[str, BarisCoretax] = {}
    for c in coretax:
        tujuan = batal if c.batal else valid
        if c.ntpn in tujuan:
            lama = tujuan[c.ntpn]
            lama.amount += c.amount
            lama.sisa = None if lama.sisa is None or c.sisa is None else lama.sisa + c.sisa
            lama.jumlah_baris += 1
        else:
            tujuan[c.ntpn] = dataclasses.replace(c)
    for c in valid.values():
        if c.jumlah_baris > 1:
            catatan.append(f"NTPN {c.ntpn} muncul {c.jumlah_baris} kali di Coretax; nilainya dijumlahkan ({rupiah(c.amount)}).")
    if batal:
        catatan.append(f"{len(batal)} NTPN berstatus dibatalkan (IsCancelled) di Coretax tidak dihitung sebagai NTPN valid.")

    dipakai: dict[str, BarisRekon] = {}
    for r in rekon:
        t = r.terima
        if not t.ntpn:
            if _sama(t.terima, 0):
                r.status = NIHIL
                r.tindakan.append("Tidak ada potongan (Rp0) - tidak perlu tindakan")
            else:
                r.status = TIDAK_DITEMUKAN
                r.tindakan.append("NTPN kosong padahal ada potongan - isi NTPN dari bukti setor")
        elif t.ntpn in valid:
            if t.ntpn in dipakai:
                r.status = NTPN_GANDA
                r.tindakan.append(f"NTPN sama dengan BKU No {dipakai[t.ntpn].terima.no} - "
                                  "cek apakah dicatat dobel atau NTPN salah salin")
            else:
                r.ct = valid[t.ntpn]
                dipakai[t.ntpn] = r
                r.status = COCOK
        elif t.ntpn in batal:
            r.status = DIBATALKAN
            r.tindakan.append("NTPN ini berstatus dibatalkan di Coretax - cek setoran pengganti / NTPN baru")
        else:
            r.status = TIDAK_DITEMUKAN

    # NTPN yang tidak ketemu persis: cari padanan terdekat (salah ketik) di NTPN Coretax yang belum terpakai
    for r in rekon:
        t = r.terima
        if r.status != TIDAK_DITEMUKAN or not t.ntpn:
            continue
        terbaik = None
        for c in valid.values():
            if c.ntpn in dipakai or abs(len(c.ntpn) - len(t.ntpn)) > MAKS_SALAH_KETIK:
                continue
            d = jarak_edit(t.ntpn, c.ntpn)
            if d <= MAKS_SALAH_KETIK:
                skor = (d, not _sama(c.amount, t.terima), c.tanggal != t.tanggal)
                if terbaik is None or skor < terbaik[0]:
                    terbaik = (skor, c, d)
        if terbaik:
            _, c, d = terbaik
            r.ct = c
            dipakai[c.ntpn] = r
            r.status = SALAH_KETIK
            r.tindakan.append(f"Ganti NTPN di BKU ({r.baris_bku}) dari '{t.ntpn_asli}' menjadi '{c.ntpn}' (beda {d} karakter)")
        else:
            r.tindakan.append("NTPN tidak ada di Coretax - cek bukti setor (BPN); kemungkinan belum disetor, "
                              "salah NPWP, atau NTPN salah catat")

    for r in rekon:
        if r.ct is None and r.status in (TIDAK_DITEMUKAN, NTPN_GANDA):
            calon = [c.ntpn for c in valid.values() if c.ntpn not in dipakai and _sama(c.amount, r.terima.terima)]
            if calon:
                r.tindakan.append("Kemungkinan padanan di Coretax dengan nilai sama: " + ", ".join(calon[:3]))

    # Bandingkan nilai (Penerimaan BKU vs Amount Coretax) dan tanggal setor
    for r in rekon:
        if r.ct is None:
            continue
        if not _sama(r.selisih, 0):
            r.status = NILAI_BEDA if r.status == COCOK else f"{r.status} + {NILAI_BEDA}"
            r.tindakan.append(f"Nilai BKU {rupiah(r.terima.terima)} berbeda dengan Coretax {rupiah(r.ct.amount)} "
                              f"(selisih {rupiah(r.selisih)}) - sesuaikan dengan bukti setor")
        if r.tanggal_beda:
            if r.status == COCOK:
                r.status = TANGGAL_BEDA
            r.tindakan.append(f"Tanggal BKU {tgl(r.terima.tanggal)} berbeda dengan tanggal setor {tgl(r.ct.tanggal)} - "
                              "sesuaikan tanggal baris Pembayaran Pot. Pajak")

    coretax_saja = [c for c in valid.values() if c.ntpn not in dipakai]
    return rekon, list(valid.values()), coretax_saja, list(batal.values()), masalah_bku, catatan


def konsolidasikan(file_bku, file_coretax) -> Hasil:
    bku = baca_bku(file_bku)
    coretax, wp_nama, wp_npwp, dilewati, peringatan = baca_coretax(file_coretax)
    rekon, valid, saja, batal, masalah_bku, catatan = rekonsiliasi(bku, coretax)
    catatan[:0] = peringatan
    if dilewati:
        catatan.append(f"{dilewati} baris Coretax bukan setoran (ReferenceType selain Payment) diabaikan.")
    terima = [b for b in bku if b.jenis == "terima"]
    bayar = [b for b in bku if b.jenis == "bayar"]
    return Hasil(
        rekon=rekon, coretax_valid=valid, coretax_saja=saja, coretax_batal=batal,
        masalah_bku=masalah_bku, catatan=catatan,
        jumlah_terima=len(terima), jumlah_bayar=len(bayar),
        total_terima=round(sum(b.terima for b in terima), 2),
        total_bayar=round(sum(b.bayar for b in bayar), 2),
        total_coretax=round(sum(c.amount for c in valid), 2),
        wp_nama=wp_nama, wp_npwp=wp_npwp,
        masa=_label_masa([b.tanggal for b in terima if b.tanggal]),
        file_bku=Path(file_bku).name, file_coretax=Path(file_coretax).name,
        dibuat=dt.datetime.now())


# --------------------------------------------------------------------------- ringkasan

def ringkas_jenis_pajak(h: Hasil) -> list[dict]:
    grup: dict[str, dict] = {}

    def baru():
        return dict(n=0, bku=0.0, ct=0.0, dep_n=0, dep=0.0, lain_n=0, lain=0.0, sisa=0.0, kap=set())

    for r in h.rekon:
        g = grup.setdefault(r.terima.pajak or "(tanpa jenis pajak)", baru())
        g["n"] += 1
        g["bku"] += r.terima.terima
        if r.ct:
            g["ct"] += r.ct.amount
            g["sisa"] += r.ct.sisa or 0
            g["kap"].add(r.ct.kap_kjs)
            if r.ct.kap == KAP_DEPOSIT:
                g["dep_n"] += 1
                g["dep"] += r.ct.amount
            else:
                g["lain_n"] += 1
                g["lain"] += r.ct.amount
    if h.coretax_saja:
        g = grup.setdefault("(NTPN Coretax tidak ada di BKU)", baru())
        for c in h.coretax_saja:
            g["n"] += 1
            g["ct"] += c.amount
            g["sisa"] += c.sisa or 0
            g["kap"].add(c.kap_kjs)
            if c.kap == KAP_DEPOSIT:
                g["dep_n"] += 1
                g["dep"] += c.amount
            else:
                g["lain_n"] += 1
                g["lain"] += c.amount
    hasil = []
    for nama, g in grup.items():
        g["nama"] = nama
        g["selisih"] = round(g["bku"] - g["ct"], 2)
        g["kap"] = ", ".join(sorted(g["kap"]))
        hasil.append(g)
    return hasil


def ringkas_kap(h: Hasil) -> list[dict]:
    grup: dict[str, dict] = {}
    for c in h.coretax_valid:
        g = grup.setdefault(c.kap_kjs, dict(kap=c.kap_kjs, n=0, amount=0.0, sisa=0.0, ada_sisa=c.sisa is not None))
        g["n"] += 1
        g["amount"] += c.amount
        g["sisa"] += c.sisa or 0
    return sorted(grup.values(), key=lambda g: g["kap"])


def susun_kesimpulan(h: Hasil) -> list[tuple[str, float | None, str]]:
    """Daftar baris (teks, nilai rupiah atau None, gaya) untuk sheet Kesimpulan dan tampilan GUI."""
    L: list[tuple[str, float | None, str]] = []
    add = lambda teks, nilai=None, gaya="isi": L.append((teks, nilai, gaya))  # noqa: E731

    add("REKONSILIASI PAJAK BKU SKPD SIPD + CORETAX" + (f" - {h.wp_nama}" if h.wp_nama else ""), gaya="judul")
    add(f"NPWP {h.wp_npwp or '-'}  |  Masa {h.masa}  |  Sumber: {h.file_bku} & {h.file_coretax}  |  "
        f"Dibuat {h.dibuat:%d/%m/%Y %H:%M}", gaya="sub")
    add("")

    add("A. RINGKASAN ANGKA", gaya="bagian")
    add(f"Penerimaan Pot. Pajak di BKU ({h.jumlah_terima} baris)", h.total_terima)
    add(f"Pembayaran Pot. Pajak di BKU ({h.jumlah_bayar} baris)", h.total_bayar)
    add(f"Total Amount Buku Besar Coretax ({len(h.coretax_valid)} NTPN valid)", h.total_coretax)
    selisih = round(h.total_terima - h.total_coretax, 2)
    add("Selisih (Penerimaan BKU - Coretax)", selisih)
    if not _sama(selisih, 0):
        tanpa_padanan = sum(r.terima.terima for r in h.rekon if r.ct is None)
        beda_nilai = sum(r.selisih for r in h.rekon if r.ct is not None)
        ct_saja = sum(c.amount for c in h.coretax_saja)
        if not _sama(tanpa_padanan, 0):
            add("- dari Penerimaan BKU yang tidak punya padanan di Coretax", round(tanpa_padanan, 2), "butir")
        if not _sama(beda_nilai, 0):
            add("- dari selisih nilai pada NTPN yang berpadanan", round(beda_nilai, 2), "butir")
        if not _sama(ct_saja, 0):
            add("- dari NTPN Coretax yang tidak tercatat di BKU", round(-ct_saja, 2), "butir")
    add("")

    def jumlah(cond):
        rows = [r for r in h.rekon if cond(r)]
        return len(rows), round(sum(r.terima.terima for r in rows), 2)

    add("B. HASIL PENCOCOKAN NTPN & NILAI (Penerimaan BKU vs Amount Coretax)", gaya="bagian")
    n, v = jumlah(lambda r: r.status == COCOK)
    add(f"NTPN & nilai cocok, tanggal sama: {n} transaksi", v)
    n, v = jumlah(lambda r: r.status == TANGGAL_BEDA)
    add(f"NTPN & nilai cocok, tanggal BKU beda dengan tanggal setor: {n} transaksi", v)
    n, v = jumlah(lambda r: r.status.startswith(SALAH_KETIK))
    add(f"NTPN salah ketik di BKU (padanan ditemukan di Coretax): {n} transaksi", v)
    n, v = jumlah(lambda r: r.ct is not None and not _sama(r.selisih, 0))
    add(f"NTPN berpadanan tetapi nilai berbeda: {n} transaksi", v)
    n, v = jumlah(lambda r: r.status == TIDAK_DITEMUKAN)
    add(f"NTPN BKU tidak ditemukan di Coretax: {n} transaksi", v)
    for status, label in ((NTPN_GANDA, "NTPN ganda di BKU"), (DIBATALKAN, "NTPN dibatalkan di Coretax")):
        n, v = jumlah(lambda r, s=status: r.status == s)
        if n:
            add(f"{label}: {n} transaksi", v)
    add(f"NTPN Coretax tidak tercatat di BKU: {len(h.coretax_saja)} NTPN",
        round(sum(c.amount for c in h.coretax_saja), 2))
    n, _ = jumlah(lambda r: r.status == NIHIL)
    if n:
        add(f"Baris nihil (Rp0, tanpa NTPN) - tidak perlu tindakan: {n} transaksi")
    add("")

    add("C. YANG PERLU DIPERBAIKI", gaya="bagian")
    k = 0

    def item(teks):
        nonlocal k
        k += 1
        add(f"{k}. {teks}")

    typo = [r for r in h.rekon if r.status.startswith(SALAH_KETIK)]
    if typo:
        item(f"Koreksi {len(typo)} NTPN salah ketik di BKU (baris Penerimaan dan Pembayaran):")
        for r in typo:
            add(f"- {_awalan_bukti(r.terima)}{r.tindakan[0]} [{rupiah(r.terima.terima)}]", gaya="butir")
    beda = [r for r in h.rekon if r.ct is not None and not _sama(r.selisih, 0)]
    if beda:
        item(f"Sesuaikan nilai {len(beda)} transaksi yang Penerimaan BKU-nya berbeda dengan Amount Coretax:")
        for r in beda:
            add(f"- {_awalan_bukti(r.terima)}BKU {r.baris_bku}, NTPN {r.ct.ntpn}: BKU {rupiah(r.terima.terima)} vs Coretax "
                f"{rupiah(r.ct.amount)} (selisih {rupiah(r.selisih)})", gaya="butir")
    hilang = [r for r in h.rekon if r.status in (TIDAK_DITEMUKAN, NTPN_GANDA, DIBATALKAN)]
    if hilang:
        item(f"Telusuri {len(hilang)} transaksi BKU yang NTPN-nya tidak valid di Coretax:")
        for r in hilang:
            add(f"- {_awalan_bukti(r.terima)}BKU {r.baris_bku}, NTPN '{r.terima.ntpn_asli or '-'}', "
                f"{rupiah(r.terima.terima)} [{r.status}]: {'; '.join(r.tindakan)}", gaya="butir")
    if h.coretax_saja:
        total = round(sum(c.amount for c in h.coretax_saja), 2)
        item(f"Telusuri {len(h.coretax_saja)} NTPN di Coretax yang tidak tercatat di BKU (total {rupiah(total)}). "
             "Bila potongan pajak SKPD, catat di BKU; bila salah setor/bukan milik SKPD, ajukan Pemindahbukuan (Pbk):")
        for c in h.coretax_saja:
            add(f"- NTPN {c.ntpn}, {rupiah(c.amount)}, setor {tgl(c.tanggal)}, KAP {c.kap_kjs}", gaya="butir")
    if h.masalah_bku:
        item(f"Perbaiki {len(h.masalah_bku)} ketidaksesuaian pasangan Penerimaan/Pembayaran Pot. Pajak di BKU:")
        for m in h.masalah_bku:
            add(f"- {m}", gaya="butir")
    tgl_beda = [r for r in h.rekon if r.tanggal_beda]
    if tgl_beda:
        item(f"Sesuaikan tanggal {len(tgl_beda)} transaksi yang tanggal BKU-nya berbeda dengan tanggal setor (NTPN) di "
             "Coretax; tanggal baris 'Pembayaran Pot. Pajak' sebaiknya mengikuti tanggal NTPN:")
        for r in tgl_beda:
            add(f"- {_awalan_bukti(r.terima)}BKU {r.baris_bku}, NTPN {r.ct.ntpn}: tanggal BKU {tgl(r.terima.tanggal)}, "
                f"tanggal setor {tgl(r.ct.tanggal)} [{rupiah(r.terima.terima)}]", gaya="butir")
    deposit = [c for c in h.coretax_valid if c.kap == KAP_DEPOSIT and c.sisa]
    if deposit:
        item(f"{len(deposit)} setoran KAP {KAP_DEPOSIT} (deposit) masih menyisakan saldo "
             f"{rupiah(round(sum(c.sisa for c in deposit), 2))} yang belum dipakai (AmountLeft > 0). "
             "Alokasikan saat membuat bukti potong / SPT Masa (PPh 21, PPh Unifikasi, PPN) sebelum batas waktu lapor; "
             "bila tidak, pajak tercatat kurang bayar dan deposit menggantung.")
    campur = [g for g in ringkas_jenis_pajak(h) if not g["nama"].startswith("(") and g["dep_n"] and g["lain_n"]]
    for g in campur:
        item(f"{g['nama']} disetor dengan dua cara: {g['dep_n']} transaksi lewat deposit {KAP_DEPOSIT} "
             f"({rupiah(round(g['dep'], 2))}) dan {g['lain_n']} transaksi lewat KAP lain ({rupiah(round(g['lain'], 2))}). "
             "Pastikan yang lewat deposit ikut dikreditkan saat pelaporan, dan seragamkan cara setor ke depannya.")
    if k == 0:
        add("Tidak ada yang perlu diperbaiki - seluruh NTPN dan nilai BKU cocok dengan Buku Besar Coretax.")
    add("")

    add("D. CATATAN", gaya="bagian")
    add("- NTPN di BKU dicocokkan dengan ReferenceNumber di Coretax. NTPN yang tidak cocok persis dicari padanannya "
        f"di Coretax dengan beda paling banyak {MAKS_SALAH_KETIK} karakter.", gaya="butir")
    if not h.masalah_bku and _sama(h.total_terima, h.total_bayar):
        add("- Penerimaan dan Pembayaran Pot. Pajak di BKU seimbang; NTPN, ID Billing dan nilai pasangan barisnya konsisten.",
            gaya="butir")
    for c in h.catatan:
        add(f"- {c}", gaya="butir")
    return L


# --------------------------------------------------------------------------- tulis Excel

FMT_RP = "#,##0;[Red]-#,##0"
FMT_TGL = "dd/mm/yyyy"
GARIS = Side(style="thin", color="BFBFBF")
BORDER = Border(left=GARIS, right=GARIS, top=GARIS, bottom=GARIS)
FONT_HEADER = Font(bold=True, color="FFFFFF")
FILL_HEADER = PatternFill("solid", fgColor="1F3A6B")
FILL_BAGIAN = PatternFill("solid", fgColor="DCE6F2")
FILL_MERAH = PatternFill("solid", fgColor="FFC7CE")
FILL_STATUS = {
    COCOK: PatternFill("solid", fgColor="C6EFCE"),
    TANGGAL_BEDA: PatternFill("solid", fgColor="FFEB9C"),
    NIHIL: PatternFill("solid", fgColor="EDEDED"),
}


def _tulis_tabel(ws, kolom, data, baris=1, warna=None) -> int:
    """kolom: list (judul, fungsi ambil nilai, jenis). jenis: teks/panjang/uang/tgl/int."""
    for j, (judul, _, _) in enumerate(kolom, 1):
        c = ws.cell(baris, j, judul)
        c.font, c.fill, c.border = FONT_HEADER, FILL_HEADER, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    lebar = [min(len(k[0]), 18) for k in kolom]
    for i, item in enumerate(data, baris + 1):
        isi = warna(item) if warna else None
        for j, (_, ambil, jenis) in enumerate(kolom, 1):
            v = ambil(item)
            c = ws.cell(i, j, v)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=jenis == "panjang")
            if jenis == "uang":
                c.number_format = FMT_RP
            elif jenis == "tgl" and isinstance(v, dt.date):
                c.number_format = FMT_TGL
            elif jenis in ("teks", "panjang"):
                c.number_format = "@"
            if isi:
                c.fill = isi
            n = 0 if v is None else len(rupiah(v)) if jenis == "uang" else 10 if jenis == "tgl" else len(str(v))
            lebar[j - 1] = max(lebar[j - 1], n)
    for j, (_, _, jenis) in enumerate(kolom, 1):
        huruf = get_column_letter(j)
        batas = 60 if jenis == "panjang" else 45
        ws.column_dimensions[huruf].width = max(ws.column_dimensions[huruf].width or 0, min(lebar[j - 1] + 2, batas))
    return baris + len(data)


def _sheet_kesimpulan(ws, baris):
    ws.column_dimensions["A"].width = 125
    ws.column_dimensions["B"].width = 20
    for i, (teks, nilai, gaya) in enumerate(baris, 1):
        a = ws.cell(i, 1, teks)
        a.alignment = Alignment(wrap_text=True, vertical="top", indent=2 if gaya == "butir" else 0)
        if nilai is not None:
            b = ws.cell(i, 2, nilai)
            b.number_format = FMT_RP
            b.alignment = Alignment(vertical="top")
        if gaya == "judul":
            a.font = Font(bold=True, size=14, color="1F3A6B")
        elif gaya == "sub":
            a.font = Font(italic=True, color="666666")
        elif gaya == "bagian":
            a.font = Font(bold=True, color="1F3A6B")
            a.fill = ws.cell(i, 2).fill = FILL_BAGIAN


def simpan_excel(h: Hasil, path) -> Path:
    path = Path(path)
    wb = Workbook()

    ws = wb.active
    ws.title = "Kesimpulan"
    _sheet_kesimpulan(ws, susun_kesimpulan(h))

    kolom_rekon = [
        ("No BKU (Terima)", lambda r: r.terima.no, "int"),
        ("No BKU (Bayar)", lambda r: r.bayar.no if r.bayar else None, "int"),
        ("Tanggal BKU", lambda r: r.terima.tanggal or r.terima.tanggal_teks, "tgl"),
        ("Nomor Bukti", lambda r: r.terima.bukti, "teks"),
        ("Uraian", lambda r: r.terima.uraian, "panjang"),
        ("Jenis Pajak", lambda r: r.terima.pajak, "teks"),
        ("ID Billing", lambda r: r.terima.billing, "teks"),
        ("NTPN BKU", lambda r: r.terima.ntpn_asli, "teks"),
        ("Penerimaan BKU", lambda r: r.terima.terima, "uang"),
        ("NTPN Coretax", lambda r: r.ct.ntpn if r.ct else None, "teks"),
        ("Amount Coretax", lambda r: r.ct.amount if r.ct else None, "uang"),
        ("Selisih (BKU - Coretax)", lambda r: r.selisih, "uang"),
        ("KAP-KJS", lambda r: r.ct.kap_kjs if r.ct else None, "teks"),
        ("Tgl Setor Coretax", lambda r: r.ct.tanggal if r.ct else None, "tgl"),
        ("Saldo Belum Terpakai (AmountLeft)", lambda r: r.ct.sisa if r.ct else None, "uang"),
        ("Status", lambda r: r.status, "teks"),
        ("Tindakan Perbaikan", lambda r: "; ".join(r.tindakan), "panjang"),
    ]
    warna_status = lambda r: FILL_STATUS.get(r.status, FILL_MERAH)  # noqa: E731

    ws = wb.create_sheet("Rekonsiliasi NTPN")
    akhir = _tulis_tabel(ws, kolom_rekon, h.rekon, warna=warna_status)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom_rekon))}{akhir}"

    ws = wb.create_sheet("Perlu Perbaikan")
    perbaikan = sorted(h.perlu_perbaikan, key=lambda r: (r.status == TANGGAL_BEDA, r.terima.no))
    if perbaikan:
        pilih = ("No BKU (Terima)", "No BKU (Bayar)", "Tanggal BKU", "Uraian", "Jenis Pajak", "NTPN BKU",
                 "NTPN Coretax", "Penerimaan BKU", "Amount Coretax", "Selisih (BKU - Coretax)",
                 "Tgl Setor Coretax", "Status", "Tindakan Perbaikan")
        kolom = [k for k in kolom_rekon if k[0] in pilih]
        akhir = _tulis_tabel(ws, kolom, perbaikan, warna=warna_status)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{akhir}"
    else:
        ws["A1"] = "Tidak ada transaksi BKU yang perlu diperbaiki."

    ws = wb.create_sheet("Coretax Tidak di BKU")
    if h.coretax_saja:
        _tulis_tabel(ws, [
            ("NTPN", lambda c: c.ntpn, "teks"),
            ("Amount", lambda c: c.amount, "uang"),
            ("Saldo Belum Terpakai (AmountLeft)", lambda c: c.sisa, "uang"),
            ("KAP-KJS", lambda c: c.kap_kjs, "teks"),
            ("Masa", lambda c: c.masa, "teks"),
            ("Tgl Setor", lambda c: c.tanggal, "tgl"),
            ("Posting", lambda c: c.posting, "teks"),
            ("Tindakan Perbaikan", lambda c: "Telusuri bukti setor: bila potongan pajak SKPD, catat di BKU "
                                             "(Penerimaan & Pembayaran Pot. Pajak); bila salah setor, ajukan "
                                             "Pemindahbukuan (Pbk)", "panjang"),
        ], h.coretax_saja, warna=lambda c: FILL_MERAH)
    else:
        ws["A1"] = "Semua NTPN di Buku Besar Coretax sudah tercatat di BKU."

    ws = wb.create_sheet("Konsolidasi Jenis Pajak")
    kolom_jenis = [
        ("Jenis Pajak", lambda g: g["nama"], "teks"),
        ("Jumlah Transaksi", lambda g: g["n"], "int"),
        ("Total BKU", lambda g: round(g["bku"], 2), "uang"),
        ("Total Coretax", lambda g: round(g["ct"], 2), "uang"),
        ("Selisih (BKU - Coretax)", lambda g: g["selisih"], "uang"),
        (f"Via Deposit {KAP_DEPOSIT} (trx)", lambda g: g["dep_n"], "int"),
        (f"Via Deposit {KAP_DEPOSIT} (Rp)", lambda g: round(g["dep"], 2), "uang"),
        ("Via KAP Lain (trx)", lambda g: g["lain_n"], "int"),
        ("Via KAP Lain (Rp)", lambda g: round(g["lain"], 2), "uang"),
        ("Saldo Belum Terpakai (AmountLeft)", lambda g: round(g["sisa"], 2), "uang"),
        ("KAP-KJS Dipakai", lambda g: g["kap"], "teks"),
    ]
    jenis = ringkas_jenis_pajak(h)
    akhir = _tulis_tabel(ws, kolom_jenis, jenis)
    _baris_total(ws, kolom_jenis, 2, akhir)

    mulai = akhir + 4
    ws.cell(mulai - 1, 1, "Ringkasan Buku Besar Coretax per KAP-KJS").font = Font(bold=True, color="1F3A6B")
    kolom_kap = [
        ("KAP-KJS", lambda g: g["kap"] + (" (deposit)" if g["kap"].startswith(KAP_DEPOSIT) else ""), "teks"),
        ("Jumlah NTPN", lambda g: g["n"], "int"),
        ("Total Amount", lambda g: round(g["amount"], 2), "uang"),
        ("Sudah Terpakai", lambda g: round(g["amount"] - g["sisa"], 2) if g["ada_sisa"] else None, "uang"),
        ("Saldo Belum Terpakai (AmountLeft)", lambda g: round(g["sisa"], 2) if g["ada_sisa"] else None, "uang"),
    ]
    akhir2 = _tulis_tabel(ws, kolom_kap, ringkas_kap(h), baris=mulai)
    _baris_total(ws, kolom_kap, mulai + 1, akhir2)

    wb.save(path)
    return path


def _baris_total(ws, kolom, dari, sampai):
    t = sampai + 1
    ws.cell(t, 1, "TOTAL")
    for j, (_, _, jenis) in enumerate(kolom, 1):
        c = ws.cell(t, j)
        if jenis in ("uang", "int") and j > 1:
            huruf = get_column_letter(j)
            c.value = f"=SUM({huruf}{dari}:{huruf}{sampai})"
            c.number_format = FMT_RP if jenis == "uang" else "0"
        c.font = Font(bold=True)
        c.fill = FILL_BAGIAN
        c.border = BORDER


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    sys.stdout.reconfigure(encoding="utf-8")
    hasil = konsolidasikan(sys.argv[1], sys.argv[2])
    simpan_excel(hasil, sys.argv[3])
    for teks, nilai, _ in susun_kesimpulan(hasil):
        print(teks + (f"  {rupiah(nilai)}" if nilai is not None else ""))
    print(f"\nTersimpan: {sys.argv[3]}")
